import shutil
from datetime import datetime, timezone
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook

app = FastAPI()

app.mount("/static", StaticFiles(directory="../frontend"), name="static")
templates = Jinja2Templates(directory="../templates")

UPLOAD_DIR = Path(__file__).parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
posts: list[dict] = [
    {
        "id": 1,
        "author": "Corey Schafer",
        "title": "FastAPI is Awesome",
        "content": "This framework is really easy to use and super fast.",
        "date_posted": "April 20, 2025",
    },
    {
        "id": 2,
        "author": "Jane Doe",
        "title": "Python is Great for Web Development",
        "content": "Python is a great language for web development, and FastAPI makes it even better.",
        "date_posted": "April 21, 2025",
    },
    {
        "id": 3,
        "author": "Alex Rivera",
        "title": "Getting Started with Pydantic",
        "content": "Pydantic makes data validation in Python clean and simple. Here's how to get started.",
        "date_posted": "April 22, 2025",
    },
    {
        "id": 4,
        "author": "Maria Chen",
        "title": "SQLAlchemy and FastAPI: A Perfect Pair",
        "content": "Combining SQLAlchemy with FastAPI gives you a powerful and flexible backend stack.",
        "date_posted": "April 23, 2025",
    },
    {
        "id": 5,
        "author": "James Okafor",
        "title": "Async Programming in Python",
        "content": "Understanding async and await in Python is key to building high-performance APIs.",
        "date_posted": "April 24, 2025",
    },
    {
        "id": 6,
        "author": "Priya Nair",
        "title": "Deploying FastAPI with Docker",
        "content": "Containerizing your FastAPI app with Docker makes deployment consistent and reproducible.",
        "date_posted": "April 25, 2025",
    },
]

@app.get("/", include_in_schema=False, name = "home")
@app.get("/posts", include_in_schema=False, name="posts")
def home(request: Request):
    return templates.TemplateResponse(request, "home.html", {"posts": posts, "title": "Home"})

@app.get("/api/posts")
def get_posts():
    return posts


@app.post("/api/upload")
async def upload_excel(file: UploadFile = File(...)):
    safe_name = Path(file.filename).name
    if Path(safe_name).suffix.lower() != ".xlsx":
        raise HTTPException(status_code=400, detail="Only .xlsx files are accepted.")

    destination = UPLOAD_DIR / safe_name
    with destination.open("wb") as out_file:
        shutil.copyfileobj(file.file, out_file)

    try:
        workbook = load_workbook(destination, read_only=True, data_only=True)
    except Exception as exc:
        destination.unlink(missing_ok=True)
        raise HTTPException(
            status_code=400, detail="Could not read this file as an Excel workbook."
        ) from exc

    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1), [])
    headers = [cell.value for cell in header_row if cell.value is not None]
    row_count = max(sheet.max_row - 1, 0)
    column_count = sheet.max_column
    workbook.close()

    return {
        "filename": safe_name,
        "sheet_name": sheet.title,
        "rows": row_count,
        "columns": column_count,
        "headers": headers,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
    }